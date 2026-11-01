#!/usr/bin/env python3
"""Usiphus automated specialty deck generator.

Generates one canonical card set and exports it to Anki (.apkg) and PowerPoint
(.pptx). OpenAI generation is optional and explicitly invokes a billable API.
No API call is made unless --generate is selected and OPENAI_API_KEY is set.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

from usiphus_core import coerce_record, evaluate_quality, report_to_dict
from usiphus_v7 import export_apkg

API_URL = "https://api.openai.com/v1/chat/completions"
DEFAULT_MODEL = "gpt-4.1-mini"
MAX_RESPONSE_BYTES = 2_000_000


def load_json(path: Path) -> Any:
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError(f"Input too large: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _extract_json(text: str) -> Any:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^\```(?:json)?\s*|\s*\```$", "", text, flags=re.I)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < start:
        raise ValueError("Model response did not contain a JSON object")
    return json.loads(text[start:end + 1])


def generate_topic_cards(topic: dict[str, Any], *, specialty: str, level: str,
                         model: str, api_key: str, count: int) -> list[dict[str, Any]]:
    objectives = topic.get("learning_objectives", [])
    source_refs = topic.get("reference_sources", [])
    payload = {
        "specialty": specialty, "learner_level": level,
        "topic": topic.get("name") or topic.get("topic"),
        "learning_objectives": objectives,
        "faculty_supplied_reference_sources": source_refs,
        "requested_card_count": count,
    }
    system = (
        "You are an educational drafting assistant for medical faculty. "
        "Create accurate, appropriately qualified educational drafts, not clinical advice. "
        "Do not invent references, DOI values, guideline dates, statistics, or quotations. "
        "Use only faculty-supplied references; if none are supplied, return references as []. "
        "Flag uncertain or guideline-dependent claims in the answer for faculty verification. "
        "Return ONLY a JSON object with a cards array. Each card must have card_id, "
        "specialty, card_type (basic, mcq, viva, osce, clinical_case), difficulty, "
        "prompt, answer, options, correct_option, learning_objective, tags, references, "
        "cognitive_level, exam_domain, rubric, evidence_status. Exactly one correct MCQ option. "
        "Use original wording; do not copy textbook passages. evidence_status must be unreviewed."
    )
    request_body = {
        "model": model,
        "temperature": 0.2,
        "max_tokens": 5000,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
    }
    req = urllib.request.Request(
        API_URL,
        data=json.dumps(request_body).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=90) as response:
            raw = response.read(MAX_RESPONSE_BYTES + 1)
    except urllib.error.HTTPError as exc:
        detail = exc.read(4000).decode("utf-8", "replace")
        raise RuntimeError(f"OpenAI API HTTP {exc.code}: {detail}") from exc
    if len(raw) > MAX_RESPONSE_BYTES:
        raise ValueError("API response exceeded safety size limit")
    result = json.loads(raw)
    content = result["choices"][0]["message"]["content"]
    parsed = _extract_json(content)
    cards = parsed.get("cards")
    if not isinstance(cards, list) or not cards:
        raise ValueError("Model response must contain a non-empty cards array")
    prefix = re.sub(r"[^A-Za-z0-9]+", "-", str(topic.get("name") or topic.get("topic") or "topic")).strip("-").upper()
    for index, card in enumerate(cards, 1):
        if not isinstance(card, dict):
            raise ValueError("Every generated card must be an object")
        card.setdefault("card_id", f"AI-{prefix[:24]}-{index:03d}")
        card["specialty"] = specialty
        card["evidence_status"] = "unreviewed"
        card.setdefault("source", "AI-generated draft; faculty verification required")
    return cards


def load_curriculum(path: Path) -> dict[str, Any]:
    data = load_json(path)
    if not isinstance(data, dict) or not isinstance(data.get("topics"), list) or not data["topics"]:
        raise ValueError("Curriculum JSON must be an object with a non-empty topics array")
    if not data.get("specialty"):
        raise ValueError("Curriculum must specify specialty")
    for item in data["topics"]:
        if not isinstance(item, dict) or not (item.get("name") or item.get("topic")):
            raise ValueError("Each topic needs a name")
    return data


def export_powerpoint(cards: list[Any], destination: Path, title: str, level: str) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PowerPoint export requires optional dependency: pip install python-pptx") from exc

    prs = Presentation()
    prs.core_properties.title = title
    prs.core_properties.subject = f"{level} medical education; faculty-review draft"
    prs.core_properties.author = "Usiphus"
    blank = prs.slide_layouts[6]

    def add_text(slide: Any, text: str, x: float, y: float, w: float, h: float,
                 size: int, *, bold: bool = False) -> None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = Inches(0.12)
        frame.margin_right = Inches(0.12)
        p = frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(31, 48, 68)

    # Opening slide with visible governance status.
    slide = prs.slides.add_slide(blank)
    add_text(slide, title, 0.7, 1.0, 12.0, 1.3, 30, bold=True)
    add_text(slide, f"{level} • {len(cards)} learning cards", 0.8, 2.5, 11.5, 0.6, 20)
    add_text(slide, "AI-assisted draft — not clinically approved. Faculty must verify every claim and reference.",
             0.8, 4.7, 11.5, 1.0, 17, bold=True)

    for card in cards:
        slide = prs.slides.add_slide(blank)
        add_text(slide, card.prompt, 0.55, 0.35, 12.1, 1.15, 23, bold=True)
        y = 1.7
        if card.options:
            add_text(slide, "\n".join(f"{chr(65+i)}. {v}" for i, v in enumerate(card.options)),
                     0.8, y, 11.7, min(2.0, 0.42 * len(card.options) + 0.2), 17)
            y += min(2.2, 0.42 * len(card.options) + 0.35)
        add_text(slide, "Answer / teaching points", 0.8, y, 11.7, 0.35, 16, bold=True)
        add_text(slide, card.answer, 0.8, y + 0.42, 11.7, 2.0, 16)
        footer = f"{card.card_id} • {card.card_type} • {card.evidence_status}"
        if card.references:
            footer += " • " + "; ".join(card.references)
        add_text(slide, footer[:500], 0.8, 6.8, 11.7, 0.35, 9)
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destination)


def build(args: argparse.Namespace) -> int:
    if args.generate and args.cards:
        raise ValueError("Choose --generate with --curriculum, or --cards for existing cards")
    raw_cards: list[dict[str, Any]] = []
    level = args.level
    specialty = args.specialty
    if args.generate:
        if not args.curriculum:
            raise ValueError("--curriculum is required with --generate")
        api_key = os.environ.get("OPENAI_API_KEY", "")
        if not api_key:
            raise ValueError("Set OPENAI_API_KEY; no API key is stored in files or source code")
        curriculum = load_curriculum(Path(args.curriculum))
        specialty = specialty or curriculum["specialty"]
        level = level or curriculum.get("learner_level", "medical students")
        for topic in curriculum["topics"]:
            count = int(topic.get("card_count", args.cards_per_topic))
            if count < 1 or count > 30:
                raise ValueError("Each topic card_count must be between 1 and 30")
            raw_cards.extend(generate_topic_cards(
                topic, specialty=specialty, level=level, model=args.model,
                api_key=api_key, count=count,
            ))
    elif args.cards:
        payload = load_json(Path(args.cards))
        raw_cards = payload.get("cards", payload) if isinstance(payload, dict) else payload
        if not isinstance(raw_cards, list) or not raw_cards:
            raise ValueError("Cards file must contain a non-empty list")
        specialty = specialty or "ophthalmology"
        level = level or "medical students"
    else:
        raise ValueError("Supply --generate --curriculum or --cards")

    cards = [coerce_record(item, default_specialty=specialty, source_label="Usiphus automated generator", index=i)
             for i, item in enumerate(raw_cards)]
    report = evaluate_quality(cards, require_reviewed=args.require_reviewed)
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    (out / "generated_cards.json").write_text(
        json.dumps({"specialty": specialty, "learner_level": level,
                    "generation_notice": "AI-generated draft; faculty review required",
                    "cards": raw_cards}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (out / "quality_report.json").write_text(
        json.dumps({**report_to_dict(report), "faculty_review_required": True,
                    "clinical_accuracy_verified": False}, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8")
    if report.error_count and not args.force:
        print(f"Quality gate blocked export: {report.error_count} errors; inspect {out / 'quality_report.json'}",
              file=sys.stderr)
        return 2
    if args.formats in {"both", "anki"}:
        export_apkg(cards, out / "ophthalmology.apkg", args.deck_name or f"Usiphus::{specialty.title()}")
    if args.formats in {"both", "powerpoint"}:
        export_powerpoint(cards, out / "ophthalmology.pptx",
                          args.deck_title or f"{specialty.title()} — Medical Education", level)
    print(f"Generated {len(cards)} cards in {out.resolve()}")
    print(f"Quality: {report.error_count} errors, {report.warning_count} warnings")
    return 0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    source = p.add_mutually_exclusive_group(required=True)
    source.add_argument("--generate", action="store_true", help="Generate content with the OpenAI API")
    source.add_argument("--cards", help="Use an existing Usiphus card JSON file")
    p.add_argument("--curriculum", help="Curriculum JSON with specialty and topics")
    p.add_argument("--specialty", default="ophthalmology")
    p.add_argument("--level", default="medical students")
    p.add_argument("--model", default=os.getenv("USIPHUS_OPENAI_MODEL", DEFAULT_MODEL))
    p.add_argument("--cards-per-topic", type=int, default=6)
    p.add_argument("--formats", choices=("both", "anki", "powerpoint"), default="both")
    p.add_argument("--output-dir", default="build/ophthalmology")
    p.add_argument("--deck-name")
    p.add_argument("--deck-title")
    p.add_argument("--require-reviewed", action="store_true")
    p.add_argument("--force", action="store_true", help="Explicitly override structural quality errors")
    return p.parse_args(argv)


if __name__ == "__main__":
    try:
        raise SystemExit(build(parse_args()))
    except (ValueError, RuntimeError, OSError, KeyError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
